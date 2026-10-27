from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
# Some hosted build runners hide user/site package locations from the Python process.
# Install document dependencies into a temporary target directory and add that exact
# directory to this process's import path, so the build uses the same libraries it installed.
_DEPENDENCIES_TMP = tempfile.TemporaryDirectory(prefix='sipandu-doc-deps-')
try:
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    from PIL import Image
    from docx import Document
except ImportError:
    subprocess.check_call([
        sys.executable, '-m', 'pip', 'install', '--disable-pip-version-check',
        '--upgrade', '--target', _DEPENDENCIES_TMP.name,
        'PyMuPDF', 'Pillow', 'python-docx'
    ])
    sys.path.insert(0, _DEPENDENCIES_TMP.name)
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    from PIL import Image
    from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
UPLOADS = ROOT / 'public' / 'uploads'
SOURCE_PDF = UPLOADS / 'Surat Tugas.pdf'
OUT = UPLOADS / 'Surat Tugas - Template Editable Terkunci.docx'
UPLOADS.mkdir(parents=True, exist_ok=True)
if not SOURCE_PDF.exists():
    raise FileNotFoundError(f'Surat Tugas source PDF is missing: {SOURCE_PDF}')

# Crop the actual source PDF so the official letterhead and electronic-signature QR
# are faithful to the uploaded original and cannot be edited inside the template.
_ASSET_TMP = tempfile.TemporaryDirectory(prefix='sipandu-surat-tugas-')
_asset_dir = Path(_ASSET_TMP.name)
with fitz.open(SOURCE_PDF) as _pdf:
    _page = _pdf[0]
    _pix = _page.get_pixmap(matrix=fitz.Matrix(300 / 72, 300 / 72), alpha=False)
    _rendered = Image.frombytes('RGB', (_pix.width, _pix.height), _pix.samples)
_rendered.save(UPLOADS / 'surat-tugas-original-background.png', optimize=True)
_header = _rendered.crop((190, 105, 2290, 445)).resize((1600, 259), Image.Resampling.LANCZOS)
_header = _header.quantize(colors=48, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE)
HEADER_IMG = str(UPLOADS / 'surat-tugas-header-locked.png')
_header.save(HEADER_IMG, optimize=True)
QR_IMG = str(UPLOADS / 'surat-tugas-barcode-locked.png')
_rendered.crop((1375, 2518, 1630, 2795)).save(QR_IMG, optimize=True)

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(0.7)
sec.bottom_margin = Cm(0.65)
sec.left_margin = Cm(1.75)
sec.right_margin = Cm(1.75)
sec.header_distance = Cm(0.3)
sec.footer_distance = Cm(0.3)

# Default text style mimics the source letter.
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.line_spacing = 1.0


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV', 'start', 'end'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'nil')


def set_table_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)
            tcPr = row.cells[idx]._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(int(width * 567)))
            tcW.set(qn('w:type'), 'dxa')


def set_run_font(run, size=11, bold=False, underline=False, italic=False):
    run.font.name = 'Times New Roman'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    run.italic = italic
    return run


def add_text_paragraph(container, text='', *, align=None, before=0, after=0, line=1.0, size=11, bold=False, underline=False, italic=False, left=0, first=0, keep=False):
    p = container.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.left_indent = Cm(left) if left else None
    pf.first_line_indent = Cm(first) if first else None
    pf.keep_together = keep
    set_run_font(p.add_run(text), size=size, bold=bold, underline=underline, italic=italic)
    return p


def add_locked_content_control(paragraph, tag, alias):
    """Wrap a whole block-level paragraph in a locked Word content control."""
    p = paragraph._p
    parent = p.getparent()
    idx = parent.index(p)
    sdt = OxmlElement('w:sdt')
    sdtPr = OxmlElement('w:sdtPr')
    alias_el = OxmlElement('w:alias')
    alias_el.set(qn('w:val'), alias)
    tag_el = OxmlElement('w:tag')
    tag_el.set(qn('w:val'), tag)
    lock_el = OxmlElement('w:lock')
    lock_el.set(qn('w:val'), 'sdtContentLocked')
    id_el = OxmlElement('w:id')
    id_el.set(qn('w:val'), str(abs(hash(tag)) % 2000000000 + 1))
    sdtPr.append(alias_el)
    sdtPr.append(tag_el)
    sdtPr.append(id_el)
    sdtPr.append(lock_el)
    sdtContent = OxmlElement('w:sdtContent')
    parent.remove(p)
    sdtContent.append(p)
    sdt.append(sdtPr)
    sdt.append(sdtContent)
    parent.insert(idx, sdt)
    return sdt


def keep_next(p, val=True):
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement('w:keepNext')
    el.set(qn('w:val'), '1' if val else '0')
    pPr.append(el)

# Locked letterhead image: logo, official header, address and horizontal rule.
header_p = doc.add_paragraph()
header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
header_p.paragraph_format.space_after = Pt(5)
header_p.paragraph_format.keep_with_next = True
header_p.add_run().add_picture(HEADER_IMG, width=Cm(17.35))
add_locked_content_control(header_p, 'locked-kop-surat-kemensos', 'Kop Surat Kemensos RI (Terkunci)')

# Editable title and number: they were not among the fields the user asked to lock.
title_p = add_text_paragraph(doc, 'SURAT TUGAS', align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=0, size=14, bold=True, underline=True, keep=True)
number_p = add_text_paragraph(doc, 'NOMOR : 985/4.21/RH.00.01/10/2026', align=WD_ALIGN_PARAGRAPH.CENTER, after=13, size=11.5, bold=True, keep=True)

# Main text in an editable three-column label / colon / content table.
body = doc.add_table(rows=0, cols=3)
body.alignment = WD_TABLE_ALIGNMENT.CENTER
body.autofit = False
remove_table_borders(body)
widths = [2.3, 0.35, 14.6]
for col_index, col_width in enumerate(widths):
    body.columns[col_index].width = Cm(col_width)
# Explicit fixed table layout prevents Word/LibreOffice from distributing columns evenly.
tbl_layout = OxmlElement('w:tblLayout')
tbl_layout.set(qn('w:type'), 'fixed')
body._tbl.tblPr.append(tbl_layout)


def add_body_row(label, texts, *, top=0, bottom=3, label_bold=False, content_size=11):
    row = body.add_row()
    for i, w in enumerate(widths):
        row.cells[i].width = Cm(w)
        row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(row.cells[i], top=0, start=0, bottom=0, end=40 if i == 2 else 0)
    p0 = row.cells[0].paragraphs[0]
    p0.paragraph_format.space_before = Pt(top)
    p0.paragraph_format.space_after = Pt(bottom)
    p0.paragraph_format.line_spacing = 1.0
    set_run_font(p0.add_run(label), size=content_size, bold=label_bold)
    p1 = row.cells[1].paragraphs[0]
    p1.paragraph_format.space_before = Pt(top)
    p1.paragraph_format.space_after = Pt(bottom)
    set_run_font(p1.add_run(':'), size=content_size)
    cell = row.cells[2]
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(top)
    p.paragraph_format.space_after = Pt(0 if len(texts) > 1 else bottom)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run(texts[0]), size=content_size)
    for text in texts[1:]:
        para = cell.add_paragraph()
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(bottom)
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.left_indent = Cm(0.55)
        para.paragraph_format.first_line_indent = Cm(-0.35)
        set_run_font(para.add_run(text), size=content_size)
    return row

add_body_row('Menimbang', [
    'a. Bahwa akan melaksanakan kegiatan koordinasi bersama Layanan PDP HIV terkait Program Atensi Klaster ODHIV di Kota Padang Sidempuan;',
    'b. sehubungan dengan poin a, dipandang perlu menugaskan pegawai Direktorat RSKPNO untuk melaksanakan kegiatan dimaksud.',
], bottom=5)
add_body_row('Dasar', [
    'Daftar Isian Pelaksanaan Anggaran Direktorat RSKPNO Nomor : SP DIPA - 027.04.2.690589/2026 tanggal 1 Desember 2025.'
], top=2, bottom=4)

# MENUGASKAN lives inside the same editable table to preserve the required reading order.
menugaskan_row = body.add_row()
merged = menugaskan_row.cells[0].merge(menugaskan_row.cells[2])
set_cell_margins(merged, top=0, start=0, bottom=0, end=0)
merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
p_menugaskan = merged.paragraphs[0]
p_menugaskan.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_menugaskan.paragraph_format.space_before = Pt(0)
p_menugaskan.paragraph_format.space_after = Pt(7)
set_run_font(p_menugaskan.add_run('MENUGASKAN'), size=11.5, bold=True)
add_body_row('Kepada', [
    '1. Deny Arifin Martyon Sikumbang / Psikolog Klinis Ahli Pertama'
], bottom=5)
add_body_row('Untuk', [
    '1. Melaksanakan kegiatan koordinasi bersama Layanan PDP HIV terkait Program Atensi Klaster ODHIV di Kota Padang Sidempuan',
    '2. pelaksanaan kegiatan pada tanggal 21 s.d 23 Oktober 2026;',
    '3. tidak menerima gratifikasi dalam bentuk apapun dan diwajibkan untuk tetap menjaga integritas selama pelaksanaan tugas;',
    '4. melaporkan hasil pelaksanaan kegiatan kepada atasan.',
], bottom=2)

# Closing sentence remains editable.
closing = add_text_paragraph(
    doc,
    'Demikian disampaikan untuk dilaksanakan dengan sebaik-baiknya dan penuh rasa tanggung jawab.',
    before=8,
    after=125,
    size=11,
    line=1.0,
    keep=True,
)

# Signature block. Date stays editable; office/title, official barcode, and signer name are locked.
sig_table = doc.add_table(rows=1, cols=2)
sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
sig_table.autofit = False
remove_table_borders(sig_table)
set_table_widths(sig_table, [9.3, 7.95])
left_cell, right_cell = sig_table.rows[0].cells
left_cell.width = Cm(9.3)
right_cell.width = Cm(7.95)
for cell in (left_cell, right_cell):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
# Blank left cell provides the same offset as the source letter.
p_date = right_cell.paragraphs[0]
p_date.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_date.paragraph_format.space_after = Pt(3)
p_date.paragraph_format.keep_with_next = True
set_run_font(p_date.add_run('Jakarta, 20 Oktober 2026'), size=11)

p_title = right_cell.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_title.paragraph_format.space_before = Pt(0)
p_title.paragraph_format.space_after = Pt(2)
p_title.paragraph_format.line_spacing = 1.0
p_title.paragraph_format.keep_with_next = True
for idx, line in enumerate([
    'Direktur Rehabilitasi Sosial Korban',
    'Penyalahgunaan Narkotika, Psikotropika,',
    'dan Zat Adiktif Lainnya Dan Orang Dengan',
    'Human Immunodeficiency Virus',
]):
    if idx:
        p_title.add_run().add_break()
    set_run_font(p_title.add_run(line), size=10.7)
add_locked_content_control(p_title, 'locked-pejabat-penanda-tangan', 'Jabatan Pejabat Penanda Tangan (Terkunci)')

p_qr = right_cell.add_paragraph()
p_qr.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_qr.paragraph_format.left_indent = Cm(-0.8)
p_qr.paragraph_format.space_before = Pt(0)
p_qr.paragraph_format.space_after = Pt(0)
p_qr.paragraph_format.keep_with_next = True
p_qr.add_run().add_picture(QR_IMG, width=Cm(1.75), height=Cm(1.9))
add_locked_content_control(p_qr, 'locked-barcode-tanda-tangan-elektronik', 'Barcode Tanda Tangan Elektronik (Terkunci)')

p_name = right_cell.add_paragraph()
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_name.paragraph_format.left_indent = Cm(-0.8)
p_name.paragraph_format.space_before = Pt(0)
p_name.paragraph_format.space_after = Pt(0)
p_name.paragraph_format.keep_with_next = True
set_run_font(p_name.add_run('Anna Puspasari'), size=11.5, bold=True, underline=True)
add_locked_content_control(p_name, 'locked-nama-pejabat-penanda-tangan', 'Nama Pejabat Penanda Tangan (Terkunci)')

# The source's electronic-signature certification text is retained at the page bottom.
# It remains editable; only the requested official blocks above are locked.
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.paragraph_format.space_before = Pt(0)
footer_p.paragraph_format.space_after = Pt(0)
footer_p.paragraph_format.line_spacing = 1.0
set_run_font(footer_p.add_run('Dokumen ini telah ditandatangani secara elektronik yang diterbitkan oleh Balai Sertifikasi Elektronik (BSrE),'), size=7.5)
footer_p2 = sec.footer.add_paragraph()
footer_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p2.paragraph_format.space_after = Pt(0)
footer_p2.paragraph_format.line_spacing = 1.0
set_run_font(footer_p2.add_run('BSSN'), size=7.5)

# Document metadata and editing notes.
doc.core_properties.title = 'Surat Tugas - Template Editable dengan Bagian Terkunci'
doc.core_properties.subject = 'Template surat tugas; isi editable, kop dan identitas penanda tangan terkunci.'
doc.core_properties.author = 'SIPANDU'
doc.core_properties.keywords = 'Surat Tugas, Kemensos, editable, content controls, locked sections'

# Add an unobtrusive properties description for users opening this DOCX.
doc.core_properties.comments = 'Bagian terpilih dikunci dengan Word content controls; isi surat dapat diedit.'

doc.save(OUT)
_ASSET_TMP.cleanup()
_DEPENDENCIES_TMP.cleanup()
print(OUT)
